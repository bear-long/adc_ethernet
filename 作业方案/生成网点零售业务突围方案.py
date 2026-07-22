from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).with_name("零售业务4.0时代网点零售业务突围方案.docx")

BLUE = "17365D"
MID_BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
ORANGE = "ED7D31"
LIGHT_ORANGE = "FCE4D6"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
GRAY = "F2F2F2"
WHITE = "FFFFFF"
TEXT = "222222"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header(row):
    set_repeat_table_header(row)
    prevent_row_split(row)


def set_table_borders(table, color="B4C6E7", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, chinese="宋体", western="Arial", size=None, bold=None, color=None):
    run.font.name = western
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(run, size=9, color="666666")


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    for style_name, font_name, size, color, before, after in (
        ("Title", "黑体", 26, BLUE, 0, 14),
        ("Heading 1", "黑体", 17, BLUE, 14, 8),
        ("Heading 2", "黑体", 14, MID_BLUE, 10, 5),
        ("Heading 3", "黑体", 12, ORANGE, 8, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "正文首行缩进" not in doc.styles:
        body_style = doc.styles.add_style("正文首行缩进", WD_STYLE_TYPE.PARAGRAPH)
        body_style.base_style = normal
        body_style.font.name = "Arial"
        body_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        body_style.font.size = Pt(11)
        body_style.paragraph_format.first_line_indent = Pt(22)
        body_style.paragraph_format.line_spacing = 1.5
        body_style.paragraph_format.space_after = Pt(5)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("零售业务4.0时代网点零售业务突围方案")
    set_run_font(run, chinese="黑体", size=9, color="7F7F7F")
    add_page_number(section.footer.paragraphs[0])


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="正文首行缩进")
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, chinese="黑体", bold=True)
        second = p.add_run(text[len(bold_prefix):])
        set_run_font(second)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(3)
    bullet = p.add_run("• ")
    set_run_font(bullet, chinese="黑体", bold=True, color=MID_BLUE)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_numbered(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(5)
    r1 = p.add_run(f"{number}  {title}：")
    set_run_font(r1, chinese="黑体", bold=True, color=MID_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def add_note_box(doc, title, text, fill=PALE_BLUE, accent=MID_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(16.1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    title_run = p.add_run(title)
    set_run_font(title_run, chinese="黑体", size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.35
    p2.paragraph_format.space_after = Pt(0)
    text_run = p2.add_run(text)
    set_run_font(text_run, size=10.5)
    set_table_borders(table, color=accent, size="8")
    return table


def add_table(doc, headers, rows, widths=None, header_fill=BLUE, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0]
    set_repeat_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_run_font(run, chinese="黑体", size=font_size, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, "F8FBFE")
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_flow(doc):
    labels = [
        ("场景化获客", "把服务嵌入生活\n获得有授权的线索", LIGHT_BLUE, MID_BLUE),
        ("精细化管理", "画像分层与任务分派\n实现千人千面", LIGHT_ORANGE, ORANGE),
        ("财富管理转型", "从卖产品到做配置\n形成长期陪伴", LIGHT_GREEN, GREEN),
        ("数据复盘", "以结果反哺场景\n持续优化闭环", GRAY, BLUE),
    ]
    table = doc.add_table(rows=1, cols=7)
    table.autofit = False
    widths = [3.25, 0.55, 3.25, 0.55, 3.25, 0.55, 3.25]
    for i, width in enumerate(widths):
        set_cell_width(table.cell(0, i), width)
    for idx, (title, subtitle, fill, color) in enumerate(labels):
        cell_index = idx * 2
        cell = table.cell(0, cell_index)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, 130, 90, 130, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        set_run_font(run, chinese="黑体", size=11, bold=True, color=color)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_after = Pt(0)
        run2 = p2.add_run(subtitle)
        set_run_font(run2, size=8.5)
        if idx < 3:
            arrow_cell = table.cell(0, cell_index + 1)
            arrow_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p_arrow = arrow_cell.paragraphs[0]
            p_arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arrow = p_arrow.add_run("→")
            set_run_font(arrow, chinese="黑体", size=16, bold=True, color=MID_BLUE)
    set_table_borders(table, color=WHITE, size="0")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("零售业务4.0时代")
    set_run_font(run, chinese="黑体", size=20, bold=True, color=MID_BLUE)
    p2 = doc.add_paragraph(style="Title")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("网点零售业务突围方案")
    set_run_font(r2, chinese="黑体", size=28, bold=True, color=BLUE)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("——“场景获客 × 精细经营 × 顾问配置”三轮驱动计划")
    set_run_font(r3, chinese="楷体", size=14, color=ORANGE)

    doc.add_paragraph()
    line = doc.add_table(rows=1, cols=1)
    line.cell(0, 0).height = Cm(0.08)
    set_cell_shading(line.cell(0, 0), MID_BLUE)
    set_table_borders(line, color=MID_BLUE, size="0")

    for _ in range(4):
        doc.add_paragraph()
    info = doc.add_table(rows=5, cols=2)
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.autofit = False
    values = [
        ("网点名称", "________________________"),
        ("所在岗位", "零售客户经理（可按实际修改）"),
        ("姓　　名", "________________________"),
        ("实施周期", "90天（建议）"),
        ("提交日期", "________年____月____日"),
    ]
    for row, (label, value) in zip(info.rows, values):
        set_cell_width(row.cells[0], 3.2)
        set_cell_width(row.cells[1], 8.3)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p_label = row.cells[0].paragraphs[0]
        p_label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label_run = p_label.add_run(label)
        set_run_font(label_run, chinese="黑体", size=11, bold=True, color=BLUE)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, size=11)
    set_table_borders(info, color=WHITE, size="0")
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_heading("方案摘要", level=1)
    add_note_box(
        doc,
        "一句话方案",
        "以网点周边“三公里幸福生活圈”为经营阵地，用场景活动持续获得客户线索；用标签画像、分层分群和任务分派提升转化效率；用目标导向的资产配置和全周期陪伴沉淀客户价值，构成“获客—识别—经营—配置—复盘”闭环。",
    )
    add_body(
        doc,
        "本方案基于题目所给三种模式设计。课程总题目提到“五大制胜模式”，但后三张照片实际展示的是“场景化获客、精细化管理、财富管理转型”三种模式；这已经满足“至少选择3个模式”的要求。本方案不把三种模式简单并列，而是让场景化获客解决“客户从哪里来”，精细化管理解决“谁值得以什么方式经营”，财富管理转型解决“如何从一次成交走向长期价值”。",
    )
    add_body(
        doc,
        "方案以城区综合网点和零售客户经理岗位为示例，建议用90天完成一个经营闭环。所有数字均为建议目标，不冒充网点真实数据；正式提交前，应以第0周盘点的实际基线替换文中的“基线值/建议值”，并结合本行制度调整产品、渠道及合规表述。",
    )
    add_table(
        doc,
        ["维度", "90天建议目标", "管理含义"],
        [
            ("获客", "场景触达≥800人次；新增有效线索≥300户", "从“等客上门”转向“进入客户生活”"),
            ("经营", "客户标签完整率≥90%；有效线索首触达率≥95%", "从名单式营销转向画像式经营"),
            ("转化", "有效线索综合转化率≥15%；新增有效客户≥45户", "以客户需求和关系深化定义转化"),
            ("价值", "重点客户AUM较基线净增≥8%；配置客户数较基线增≥10%", "从单品销量转向资产配置与长期陪伴"),
            ("体验与风险", "活动满意度≥90%；重大投诉、违规销售、数据泄露为0", "增长必须以客户体验和合规为底线"),
        ],
        widths=[2.2, 6.2, 7.0],
    )

    doc.add_heading("一、对第一张照片作业要求的分析", level=1)
    add_table(
        doc,
        ["原要求", "真正考察点", "本方案如何回应"],
        [
            ("至少3个模式有机组合", "不是罗列概念，而是形成前后衔接的业务闭环", "三模式按“入口—中枢—价值出口”串联，并以数据复盘反哺场景"),
            ("说明选择理由与核心突破点", "模式必须针对网点真实痛点", "对应“客流不足、粗放营销、重销售轻陪伴”三类问题"),
            ("具体动作与成果评估", "要回答谁来做、怎么做、做到什么程度", "提供角色分工、90天节奏、SOP、量化指标和看板"),
            ("困难预估与应对", "方案要可执行、可纠偏、守合规", "覆盖客户、员工、数据、市场、合作方和合规风险"),
            ("结合所在网点和自身岗位", "体现岗位可控制、可协同的动作", "以零售客户经理牵头，联动网点负责人、大堂、理财及运营人员"),
        ],
        widths=[3.0, 5.2, 7.2],
    )
    add_note_box(
        doc,
        "写作判断",
        "高质量答案的核心不在于活动数量，而在于闭环：场景带来的每条线索都要获得授权、进入标签体系、形成下一步任务；每次咨询和配置结果又要反向判断哪个场景、哪个客群最值得继续投入。",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    doc.add_heading("二、网点现状、岗位定位与问题诊断", level=1)
    doc.add_heading("（一）适用场景与基本假设", level=2)
    add_body(
        doc,
        "本方案适用于周边同时存在居民社区、沿街商户、企事业单位或政务服务资源的城区综合网点。假设网点具备负责人、零售客户经理、理财经理、大堂经理和运营柜面等角色，但人力有限，需要通过固定节奏和协同作战提高人均产能。",
    )
    add_body(
        doc,
        "本人岗位定位为“客户经营组织者”：不只完成个人销售任务，还负责把外部场景线索转化为网点可持续经营的客户资产。可直接负责客户需求访谈、标签维护、任务跟进和重点客户咨询；需要协调网点负责人配置资源、大堂经理承接到店客户、理财经理输出专业方案、运营人员把控流程与合规。",
    )
    doc.add_heading("（二）三类核心痛点", level=2)
    add_numbered(doc, "1", "获客被动", "过度依赖自然到店和熟人转介绍，活动与客户日常生活脱节，线索数量不稳定。")
    add_numbered(doc, "2", "经营粗放", "名单同发、产品同推，客户标签不全、触达时点不准，员工做了很多动作却难以判断有效性。")
    add_numbered(doc, "3", "价值经营偏短期", "容易从产品出发找客户，对客户家庭目标、期限、流动性和风险承受能力理解不足，成交后陪伴弱。")

    doc.add_heading("（三）实施前必须补齐的基线", level=2)
    add_table(
        doc,
        ["基线项目", "第0周取数口径", "用途"],
        [
            ("客户基线", "有效客户数、近90天到店/线上活跃客户、联系方式有效率", "确定可经营盘子"),
            ("资产基线", "AUM、存款、理财及其他持仓按客户层级分布", "设定净增目标并避免只看销量"),
            ("客群基线", "养老、代发、Z世代、大众、财富、私行潜客数量", "选择优先客群"),
            ("过程基线", "月均触达、面谈、转介、配置建议、成交和复访数量", "识别漏斗堵点"),
            ("场景基线", "周边社区、商户、单位、政务与文化资源清单", "评估合作价值和触达能力"),
        ],
        widths=[2.6, 8.0, 4.8],
    )

    doc.add_heading("三、模式组合设计：三轮驱动、一个闭环", level=1)
    add_flow(doc)
    doc.add_heading("（一）为什么选择这三种模式", level=2)
    add_table(
        doc,
        ["模式", "照片中的方法要义", "选择理由", "核心突破点"],
        [
            ("场景化获客", "把金融服务无缝嵌入消费、社区、政务民生、文化社交等日常场景", "网点单纯等客上门难以获得稳定增量；周边生态资源可低成本复用", "从“网点等客”突破到“生活触点获客”"),
            ("精细化管理", "依托客户画像与标签进行分层、分群、精准营销和智能推荐", "客户需求差异大，人力有限，必须把时间用在正确客户和正确时点上", "从“千人一面”突破到“千人千面”"),
            ("财富管理转型", "从货架式产品销售转向买方视角的顾问式配置和全周期陪伴", "单品推销难以建立信任，也不利于AUM留存和交叉需求发现", "从“一次成交”突破到“长期伙伴”"),
        ],
        widths=[2.2, 4.8, 5.2, 3.2],
        font_size=8.8,
    )
    doc.add_heading("（二）组合逻辑", level=2)
    add_numbered(doc, "1", "场景是入口", "围绕三公里生活圈，把非金融高频需求与金融服务连接起来，形成可持续触点。")
    add_numbered(doc, "2", "画像是中枢", "所有获客动作统一进入标签、评分和任务体系，决定由谁、在何时、用何种内容跟进。")
    add_numbered(doc, "3", "配置是价值出口", "以客户目标而非产品库存为起点，完成需求澄清、风险匹配、配置建议、持续陪伴。")
    add_numbered(doc, "4", "复盘是飞轮", "按“场景—客群—人员—结果”分析投入产出，把高质量客群的共同特征反哺下一轮活动。")
    add_note_box(
        doc,
        "核心经营主张",
        "不以“办一场活动、卖一只产品”为终点，而以“新增一个可授权联系、可识别需求、可持续陪伴的有效客户关系”为最小经营成果。",
        fill=LIGHT_GREEN,
        accent=GREEN,
    )

    doc.add_heading("四、重点项目：“三公里幸福生活圈”", level=1)
    doc.add_heading("（一）场景地图与客群承接", level=2)
    add_table(
        doc,
        ["场景项目", "合作触点与活动", "优先客群", "画像标签", "后续承接"],
        [
            ("社区安心站", "联合社区开展反诈讲堂、社保/医保服务咨询、适老智能手机课堂、养老规划微沙龙", "养老客群、家庭照护者、社区居民", "年龄段、养老金、社保需求、数字使用、家庭照护", "适老服务→养老金与现金流梳理→养老规划→定期复访"),
            ("商圈惠生活", "链接便利店、餐饮、家政等商户，开展便民权益日、商户经营体检、支付安全咨询", "大众客户、商户、代发单位员工", "职业、消费偏好、经营年限、结算、融资、代发意向", "权益触达→账户活跃→结算/代发需求→家庭资产配置"),
            ("政务便民角", "围绕社保、医保、生活缴费等低频刚需，提供流程指引和金融知识咨询；不得越权代办", "养老客群、新市民、大众客户", "政务事项、缴费习惯、常住区域、服务偏好", "事项咨询→便民服务→账户使用→长期关系"),
            ("青年共创局", "开展读书、运动、亲子、职业成长等小型社群活动，设置金融知识互动而非硬推销", "Z世代、年轻代发客户、年轻家庭", "兴趣圈层、人生阶段、收支习惯、保障与住房目标", "社群互动→收支诊断→定投/保障教育→成长陪伴"),
        ],
        widths=[2.2, 5.5, 2.7, 2.9, 4.1],
        font_size=8.3,
    )
    doc.add_heading("（二）活动设计原则", level=2)
    add_bullet(doc, "生活需求在前、金融服务在后：活动内容至少70%解决真实生活问题，避免把场景活动做成产品宣讲会。")
    add_bullet(doc, "小而高频、固定节奏：每周至少1次外拓触点、每两周1次主题活动，比一次大型活动更利于测试和复用。")
    add_bullet(doc, "一场一客群、一场一目标：每场只设1个主客群和1—2个主转化动作，避免指标过多导致现场失焦。")
    add_bullet(doc, "授权留资、最小必要：明确告知信息用途和联系范围，客户自愿授权；不通过合作方私下交换名单。")
    add_bullet(doc, "可跟进才算有效：没有负责人、没有下一步日期、没有客户授权的名单，不计入有效线索。")

    doc.add_heading("（三）单场活动标准作业程序（SOP）", level=2)
    add_table(
        doc,
        ["时间", "动作", "责任人", "必留成果/标准"],
        [
            ("T-7至T-3", "确定主客群、客户痛点、合作分工、物料和风险审查；邀约名单去重", "零售客户经理牵头；负责人审批；运营复核", "一页活动卡：目标、预算、客群、话术、合规点、责任人"),
            ("T-2至T-1", "分层邀约；确认到场；准备登记授权和需求问卷", "客户经理、大堂经理", "邀约—确认—预计到场漏斗"),
            ("T日", "签到授权、主题服务、需求互动、预约下一步；现场不强卖", "全体；客户经理主讲/主持", "授权线索、需求标签、预约事项、客户反馈"),
            ("T+1（24小时内）", "录入系统、标签补全、线索分级、分派责任人", "客户经理；零售负责人抽查", "标签完整率≥90%；100%有责任人"),
            ("T+2（48小时内）", "首次个性化联系，确认需求与下一步；拒绝客户进入免打扰", "责任客户经理", "首触达率≥95%；记录结果"),
            ("T+7", "完成咨询、到店或远程服务；必要时转介理财经理", "客户经理/理财经理", "有效线索7日推进率≥60%"),
            ("T+30", "复盘转化、AUM、体验、投诉、成本和合作方质量", "网点负责人主持", "形成保留、优化、停止清单"),
        ],
        widths=[2.0, 6.1, 3.2, 4.1],
        font_size=8.5,
    )

    doc.add_heading("五、精细化管理：从线索到“一户一策”", level=1)
    doc.add_heading("（一）标签、分层与优先级", level=2)
    add_body(
        doc,
        "标签不是越多越好，而是要能触发动作。建议建立五类最小标签：基本特征（年龄、职业、家庭阶段）、场景来源、需求目标、服务偏好、风险与合规状态。客户自主提供的信息与系统已有信息应按权限使用，敏感信息不得记录在个人手机、微信群或纸质散表中。",
    )
    add_table(
        doc,
        ["优先级", "判定示例", "服务时效", "经营策略"],
        [
            ("A：明确需求", "已预约咨询；存在到期、资金安排、养老/教育/住房等明确目标", "24小时内联系，7日内形成服务方案", "客户经理主办，必要时转理财经理或1+N专家"),
            ("B：潜在需求", "标签匹配重点客群，有互动但尚无明确时点", "48小时内联系，30日内完成二次培育", "发送匹配内容，邀请小型沙龙或需求诊断"),
            ("C：一般培育", "仅完成活动互动，需求不明确或暂不接受咨询", "按客户偏好低频触达", "进入社群/内容培育，不反复电话打扰"),
            ("D：免打扰/不适配", "明确拒绝、联系方式无效、存在风险限制", "立即停止营销或按制度处理", "尊重意愿，更新状态，避免重复触达"),
        ],
        widths=[2.5, 5.1, 3.7, 4.1],
        font_size=8.7,
    )
    doc.add_heading("（二）重点客群“一户一策”", level=2)
    add_table(
        doc,
        ["客群", "照片建议策略", "网点具体动作", "员工角色"],
        [
            ("大众客户（长尾）", "线上标准化服务、智能客服", "场景扫码后推送办事指南和线上服务；复杂问题预约到店", "大堂做好引导、解释和转介"),
            ("财富客户（中产）", "专属理财顾问＋标准化产品组合", "围绕家庭应急、教育、养老等目标做资产盘点；季度回顾", "客户经理发现需求，理财经理提供专业建议"),
            ("私行潜客（高净值）", "“1+N”团队服务", "识别企业经营、税务、法律、传承等复杂需求，按权限转介专家", "不越界解答，及时调动行内专家"),
            ("代发客户", "专属定制服务、促进资金留存", "入企金融课堂、员工权益日、分层触达；关注账户活跃和家庭需求", "客户经理联动公司条线与大堂承接"),
            ("Z世代客户", "圈层认同、情感共鸣、个性表达", "读书/运动/职业成长共创；提供预算与储蓄工具教育", "做社群组织者，减少单向推销"),
            ("养老客群", "适老服务、养老规划、养老生态", "社保咨询、反诈、智能手机课堂、养老金现金流梳理", "耐心服务、积累信任、识别家庭共同决策人"),
        ],
        widths=[2.2, 3.5, 6.1, 3.6],
        font_size=8.3,
    )
    doc.add_heading("（三）客户经营“六步法”", level=2)
    add_numbered(doc, "1", "识别", "确认客户来源、授权状态、基础标签和当前事项。")
    add_numbered(doc, "2", "问需", "使用开放式问题了解人生目标、资金用途、期限、流动性和顾虑。")
    add_numbered(doc, "3", "分层", "按照需求紧迫度、服务复杂度和客户意愿确定A/B/C/D级。")
    add_numbered(doc, "4", "分派", "普通服务由大堂/客户经理承接，配置需求转理财经理，复杂需求启动1+N。")
    add_numbered(doc, "5", "跟进", "明确下一动作、责任人和完成日期，每次互动更新标签与意愿。")
    add_numbered(doc, "6", "复盘", "分析客户是否获得价值、漏斗停在哪里、下一轮场景是否需要调整。")

    doc.add_heading("六、财富管理转型：从“卖产品”到“做配置”", level=1)
    doc.add_heading("（一）顾问式服务五步流程", level=2)
    add_table(
        doc,
        ["步骤", "关键动作", "交付物", "合规底线"],
        [
            ("1. 建立信任", "说明服务流程、信息用途、人员职责；先解决客户当前问题", "服务约定与客户授权", "不诱导授权，不夸大专业能力"),
            ("2. 目标诊断", "梳理收支、负债、应急、教育、养老、购房等目标和优先级", "家庭目标清单", "只收集服务所必需的信息"),
            ("3. 风险与期限匹配", "确认风险承受能力、投资经验、资金期限、流动性约束", "风险与流动性画像", "执行本行适当性管理，不代客操作"),
            ("4. 配置建议", "按“现金管理—稳健底仓—长期增值—保障/养老”解释组合逻辑", "一页式配置建议书", "不承诺保本收益；充分揭示风险和费用"),
            ("5. 持续陪伴", "约定回顾频率；围绕目标变化、到期资金、市场波动做再评估", "月度触达、季度回顾记录", "不以短期波动诱导频繁交易"),
        ],
        widths=[2.2, 5.8, 3.3, 4.1],
        font_size=8.5,
    )
    doc.add_heading("（二）一页式配置建议书框架", level=2)
    add_bullet(doc, "客户目标：目标金额、预计时间、优先级，以及是否允许调整。")
    add_bullet(doc, "资金分桶：日常流动资金、应急资金、阶段目标资金、长期资金分别管理。")
    add_bullet(doc, "约束条件：风险承受能力、可投资期限、负债压力、流动性和家庭共同决策。")
    add_bullet(doc, "组合逻辑：说明每类资产承担什么角色、可能发生什么风险，而不是只列产品名称。")
    add_bullet(doc, "跟踪约定：下一次回顾日期、触发再平衡的条件、客户可联系的服务人员。")
    add_note_box(
        doc,
        "岗位边界",
        "客户经理负责发现需求、建立关系和组织服务；涉及具体产品建议时，应由具备相应资质的人员按本行流程办理。不得以作业方案替代本行销售制度、风险测评、双录或信息披露要求。",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    doc.add_heading("七、90天执行计划与责任分工", level=1)
    doc.add_heading("（一）四阶段推进", level=2)
    add_table(
        doc,
        ["阶段", "重点任务", "里程碑", "主要责任人"],
        [
            ("第0—2周：盘点准备", "完成基线、场景地图、合作方筛选、重点客群清单、话术与合规审查；全员演练", "形成1张地图、3份客群名单、4类活动模板、1套看板", "网点负责人/零售客户经理"),
            ("第3—6周：小步试点", "社区安心站、商圈惠生活各试点2次；严格执行24/48小时跟进", "至少4场活动；找到转化最高的2个“场景×客群”组合", "客户经理/大堂/理财经理"),
            ("第7—10周：复制放大", "保留高效活动，增加政务便民或青年共创；形成固定周节奏", "有效线索累计≥200户；A/B级客户均有下一步任务", "零售负责人统筹"),
            ("第11—13周：深耕复盘", "重点推进顾问咨询、配置回顾、老带新和合作续约；分析投入产出", "完成90天看板、案例复盘和下一周期计划", "网点负责人主持，全员参与"),
        ],
        widths=[2.7, 6.2, 4.0, 2.5],
        font_size=8.4,
    )
    doc.add_heading("（二）角色责任矩阵", level=2)
    add_table(
        doc,
        ["角色", "核心责任", "固定动作", "考核重点"],
        [
            ("网点负责人", "定目标、配资源、抓协同与合规", "周一看板会；月度合作方/风险复盘；重大事项审批", "闭环完成率、资源效率、重大风险为0"),
            ("零售客户经理（本人）", "项目牵头、场景拓展、客户分层与跟进", "每周更新场景地图；维护A/B级客户；组织活动复盘", "有效线索、7日推进、综合转化、客户体验"),
            ("理财经理", "需求诊断、配置建议、重点客户陪伴", "每周配置门诊；季度回顾；案例分享", "建议书质量、配置客户、AUM净增、适当性合规"),
            ("大堂经理", "到店识别、线上迁移、预约和转介", "每日晨会认领重点标签；当天完成转介记录", "有效转介率、客户等待时间、服务满意度"),
            ("运营/柜面", "业务办理、风险提示、异常反馈", "按流程核验；发现需求后规范转介", "差错、投诉、合规与有效转介"),
            ("合作方接口人", "资源协调、客群触达、活动共建", "活动前确认边界，活动后反馈", "到场质量、客户评价、合作稳定性"),
        ],
        widths=[2.4, 4.7, 5.4, 2.9],
        font_size=8.3,
    )
    doc.add_heading("（三）固定经营节奏", level=2)
    add_table(
        doc,
        ["频率", "会议/动作", "控制在", "输出"],
        [
            ("每日", "晨会认领当日客户任务；夕会清理未完成任务", "各10分钟", "今日3件关键事、延期原因"),
            ("每周", "周一漏斗看板；周三配置门诊；周末场景活动", "30/45/60分钟", "问题清单、客户方案、活动线索"),
            ("每两周", "场景活动复盘与AB测试", "45分钟", "保留/优化/停止清单"),
            ("每月", "经营质量与风险复盘", "60分钟", "指标、投诉、合规、成本、合作方评分"),
        ],
        widths=[2.2, 6.2, 2.6, 4.4],
    )

    doc.add_heading("八、成果评估：用一张看板管理闭环", level=1)
    doc.add_heading("（一）指标体系", level=2)
    add_table(
        doc,
        ["层级", "指标与计算口径", "90天建议值", "责任人"],
        [
            ("场景入口", "触达人次；授权线索数；有效线索率＝有效线索/授权线索", "触达≥800；有效线索≥300；有效率≥50%", "客户经理"),
            ("过程质量", "标签完整率；24小时录入率；48小时首触达率；7日推进率", "≥90%；100%；≥95%；≥60%", "客户经理/零售负责人"),
            ("客户转化", "新增有效客户；综合转化率＝形成有效客户关系/有效线索", "新增≥45户；转化率≥15%", "全体经营人员"),
            ("客户价值", "重点客户AUM净增；配置客户净增；代发活跃、账户活跃等", "AUM较基线≥8%；配置客户较基线≥10%", "理财经理/客户经理"),
            ("体验", "活动满意度；客户建议闭环率；投诉量", "满意度≥90%；建议闭环≥95%", "大堂/客户经理"),
            ("风险", "数据泄露、违规承诺、适当性违规、重大投诉", "全部为0", "网点负责人/全员"),
        ],
        widths=[2.2, 7.3, 3.6, 2.3],
        font_size=8.4,
    )
    add_body(
        doc,
        "上述目标属于示范值。若网点基线较低，应优先保证数据质量、客户体验和合规，不应通过重复留资、强制激活、短期冲量或不适当销售完成数字。最终目标应由网点负责人基于历史数据确认。",
    )
    doc.add_heading("（二）建议综合评分卡", level=2)
    add_table(
        doc,
        ["评分项", "权重", "评分方法"],
        [
            ("新增与转化", "25%", "按有效线索、新增有效客户和综合转化率完成度计分"),
            ("客户价值", "25%", "按AUM净增、配置客户净增和客户留存计分"),
            ("过程质量", "20%", "按标签、录入、首触达、7日推进和任务闭环计分"),
            ("客户体验", "15%", "按满意度、建议闭环和投诉情况计分"),
            ("能力与协同", "15%", "按案例复盘、转介质量、合作方共建和知识提升计分"),
            ("风险扣分", "一票否决/倒扣", "发生重大违规、数据泄露或严重投诉时按本行制度处理"),
        ],
        widths=[3.0, 2.0, 10.4],
    )
    doc.add_heading("（三）周看板模板", level=2)
    add_table(
        doc,
        ["指标", "本周", "累计", "目标", "达成率", "环比", "问题与下周动作"],
        [
            ("场景触达人次", "", "", "800", "", "", ""),
            ("新增有效线索", "", "", "300", "", "", ""),
            ("标签完整率", "", "", "90%", "", "", ""),
            ("48小时首触达率", "", "", "95%", "", "", ""),
            ("7日推进率", "", "", "60%", "", "", ""),
            ("新增有效客户", "", "", "45", "", "", ""),
            ("重点客户AUM净增", "", "", "8%", "", "", ""),
            ("满意度/投诉", "", "", "≥90%/0", "", "", ""),
        ],
        widths=[3.2, 1.5, 1.5, 1.5, 1.6, 1.5, 4.6],
        font_size=8.2,
    )

    doc.add_heading("九、困难预估与应对", level=1)
    add_table(
        doc,
        ["可能困难", "预警信号", "应对措施", "责任人"],
        [
            ("活动热闹但线索质量低", "到场多、授权少、需求标签空白、后续联系无回应", "缩小客群；活动前设主问题；现场预约下一步；按有效线索而非到场人数评价", "客户经理"),
            ("客户反感营销或担忧隐私", "拒绝留资、退群、投诉、要求停止联系", "明示用途与频率；最小必要采集；提供退出渠道；免打扰状态立即生效", "全员"),
            ("数据分散、标签失真", "重复客户、多人跟进、记录滞后、标签互相矛盾", "以本行合规系统为唯一主档；24小时录入；负责人每周抽查10户", "零售负责人"),
            ("员工任务多、协同弱", "线索无人认领、转介后无反馈、活动依赖少数人", "晨会认领、夕会销项；设置主办/协办；以团队闭环指标替代单纯抢客户", "网点负责人"),
            ("顾问能力不足", "仍以产品话术为主，无法解释目标、期限和风险", "每周一案例演练；理财经理配置门诊；复杂需求启动1+N专家", "理财经理"),
            ("合作方不稳定", "临时取消、过度索取客户数据、只关注补贴", "建立准入评分和书面边界；不交换客户名单；至少准备两类替代场景", "客户经理/负责人"),
            ("市场波动影响信任", "客户集中赎回、情绪化咨询、投诉上升", "主动陪伴和风险再确认；回到目标与期限；不预测短期市场、不承诺收益", "理财经理"),
            ("短期指标挤压长期经营", "只看销量、重复触达、配置偏离客户目标", "评分卡增加过程、体验、留存和风险权重；重大违规一票否决", "网点负责人"),
            ("目标与真实基线不匹配", "连续两周达成率显著偏离且非执行问题", "第0周校准目标；第4周只调资源和策略，确需调目标须留存依据并审批", "网点负责人"),
        ],
        widths=[3.0, 4.0, 6.4, 2.0],
        font_size=8.0,
    )

    doc.add_heading("十、预期成果与个人岗位提升", level=1)
    doc.add_heading("（一）网点层面", level=2)
    add_bullet(doc, "形成一套可复制的“三公里场景地图＋活动SOP＋客户看板”，减少对自然客流和临时营销的依赖。")
    add_bullet(doc, "建立从场景线索到客户画像、任务分派、顾问服务和结果复盘的全流程，提升人均经营效率。")
    add_bullet(doc, "客户评价从“卖产品的银行”向“解决生活与家庭财富问题的长期伙伴”转变。")
    add_bullet(doc, "AUM、有效客户和配置客户实现质量型增长，同时把投诉、适当性和数据安全纳入同一张看板。")
    doc.add_heading("（二）个人岗位层面", level=2)
    add_bullet(doc, "从单一产品销售者升级为场景组织者、需求发现者和客户关系经营者。")
    add_bullet(doc, "掌握客群分层、标签经营、漏斗分析和跨岗位协同，提高时间投入的精准度。")
    add_bullet(doc, "通过顾问式问需与1+N转介提升专业服务能力，在权限边界内为客户提供更完整的解决方案。")

    doc.add_heading("十一、结语", level=1)
    add_body(
        doc,
        "零售业务4.0时代，网点的竞争力不再只来自位置、产品和单次营销，而来自进入客户生活的能力、识别差异化需求的能力，以及长期陪伴客户实现家庭目标的能力。本方案以场景化获客打开增量入口，以精细化管理提升经营效率，以财富管理转型沉淀长期价值。三种模式通过统一数据、统一任务和统一评价形成闭环，既能回答“客户从哪里来”，也能回答“客户为什么留下、价值如何持续增长”。",
    )
    add_body(
        doc,
        "实施中应坚持“小步试点、数据复盘、客户自愿、合规先行”。只要每一次活动都能沉淀有效关系，每一条线索都有下一步动作，每一次配置都回到客户真实目标，网点就能够由被动服务点逐步转变为三公里生活圈内可信赖的综合金融服务伙伴。",
    )

    doc.add_page_break()
    doc.add_heading("附录：提交前个性化修改清单", level=1)
    add_table(
        doc,
        ["需替换内容", "建议填写方法", "完成"],
        [
            ("网点与岗位", "填写真实网点名称、岗位、姓名；按实际人员调整责任分工", "□"),
            ("周边资源", "把“社区/商圈/政务/青年”替换为网点周边真实合作资源名称", "□"),
            ("基线数据", "补充有效客户、AUM、客群、活动、转化等第0周数据", "□"),
            ("目标值", "根据基线与本行考核校准90天目标，保留计算口径", "□"),
            ("产品与渠道", "替换为本行合规可提供的服务，不写未经确认的产品或权益", "□"),
            ("制度表述", "按本行客户信息、营销、适当性、双录、投诉等制度复核", "□"),
            ("个人案例", "加入1个本人服务客户或组织活动的真实匿名案例，增强岗位结合度", "□"),
        ],
        widths=[3.0, 10.8, 1.6],
    )
    add_note_box(
        doc,
        "可直接用于汇报的总结",
        "我的突围思路可以概括为“三个转变”：获客上，从坐等客户转为进入生活场景；经营上，从千人一面转为分层分群、一户一策；价值创造上，从推销单品转为围绕客户目标做资产配置。最终通过90天试点，把三种模式沉淀为网点可复制、可衡量、可合规运行的经营闭环。",
        fill=LIGHT_GREEN,
        accent=GREEN,
    )

    core = doc.core_properties
    core.title = "零售业务4.0时代网点零售业务突围方案"
    core.subject = "场景化获客、精细化管理、财富管理转型三模式组合方案"
    core.author = "作业方案"
    core.keywords = "零售业务4.0, 场景化获客, 精细化管理, 财富管理"
    core.comments = "正式提交前请根据实际网点和岗位修改文中占位信息与建议目标。"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
